#!/usr/bin/env python3
"""Prueba que el corte conserva avance y continúa con cualquier actividad CMS."""

from __future__ import annotations

from datetime import datetime, timedelta
import json
from pathlib import Path
import sys
import tempfile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.build_dashboard import build_payload, load_cutover, find_header


def create_new_forms(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Respuestas Forms"
    sheet.append([
        "Id", "Hora de inicio", "Hora de finalización", "Correo electrónico",
        "Nombre", "CeCo", "Selecciona la actividad que deseas registrar", "Evidencia del avance",
    ])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        cutover_file = temp / "cutover.json"
        cutover_file.write_text(json.dumps({
            "cutoff": "2026-09-17 11:56:57",
            "baseline": str(ROOT / "cms" / "Corte_Forms_2026-09-17_115657.xlsx"),
            "label": "Corte de prueba",
        }), encoding="utf-8")
        cutover = load_cutover(cutover_file)
        assert cutover
        forms = Path(temp_dir) / "Forms_nuevo.xlsx"
        baseline = build_payload(
            ROOT / "cms" / "Sistema de Evidencias OPS.xlsx",
            ROOT / "cms" / "Directorio.xlsx",
            ROOT / "config" / "settings.json",
            ROOT / "cms" / "Sistema_Evidencias_OPS_CMS.xlsx",
            baseline_path=cutover["baseline"],
            cutoff=cutover["cutoff"],
            cutover_config_path=cutover_file,
        )
        # Usa una actividad vigente tomada del CMS. Operación puede retirar
        # Rack FHW u otra campaña sin volver obsoleta la prueba del corte.
        activity = next(
            item["name"] for item in baseline["activities"]
            if any(not store["activities"][item["name"]] for store in baseline["stores"])
        )
        candidate = next(store for store in baseline["stores"] if not store["activities"][activity])
        baseline_rows_included = baseline["quality"]["cutover"]["baselineRowsIncluded"]
        after = cutover["cutoff"] + timedelta(seconds=1)
        before = cutover["cutoff"] - timedelta(seconds=1)
        evidence = "https://grupovips-my.sharepoint.com/evidencias/prueba.jpg"
        continuation_evidence = "https://grupovips-my.sharepoint.com/evidencias/continuacion.jpg"
        create_new_forms(forms, [
            [1, after, after, "", candidate["store"], candidate["ceco"], activity, evidence],
            [2, after, after, "", candidate["store"], candidate["ceco"], "Roll Out", evidence],
            [3, before, before, "", candidate["store"], candidate["ceco"], activity, evidence],
            [4, "", "", "", candidate["store"], candidate["ceco"], activity, continuation_evidence],
        ])
        payload = build_payload(
            forms,
            ROOT / "cms" / "Directorio.xlsx",
            ROOT / "config" / "settings.json",
            ROOT / "cms" / "Sistema_Evidencias_OPS_CMS.xlsx",
            baseline_path=cutover["baseline"],
            cutoff=cutover["cutoff"],
            cutover_config_path=cutover_file,
        )
        # Retirar una actividad debe retirar también su historia y denominador.
        cms_off = temp / "cms_inactivo.xlsx"
        workbook = load_workbook(ROOT / "cms" / "Sistema_Evidencias_OPS_CMS.xlsx")
        sheet = workbook["Actividades"]
        header, columns = find_header(sheet, {"actividad", "activo", "evidencia requerida"})
        for row in range(header + 1, sheet.max_row + 1):
            if sheet.cell(row, columns["actividad"] + 1).value == activity:
                sheet.cell(row, columns["activo"] + 1, "No")
                sheet.cell(row, columns["evidencia requerida"] + 1, "No")
        workbook.save(cms_off)
        hidden = build_payload(
            forms, ROOT / "cms" / "Directorio.xlsx", ROOT / "config" / "settings.json", cms_off,
            baseline_path=cutover["baseline"], cutoff=cutover["cutoff"], cutover_config_path=cutover_file,
        )
        assert activity not in {item["name"] for item in hidden["activities"]}
        assert all(item["activity"] != activity for item in hidden["submissions"])
        assert all(activity not in item["activities"] for item in hidden["stores"])
        assert hidden["summary"]["activities"] == payload["summary"]["activities"] - 1
        cut = payload["quality"]["cutover"]
        assert cut == {
            **cut,
            "enabled": True,
            "baselineRowsIncluded": baseline_rows_included,
            "newFormsRowsRead": 4,
            "newFormsRowsIncluded": 2,
            "newFormsRowsIncludedWithoutTimestamp": 1,
            "newFormsRowsRejectedAtOrBeforeCutoff": 1,
            "newFormsRowsRejectedInactive": 1,
        }
        updated_store = next(store for store in payload["stores"] if store["ceco"] == candidate["ceco"])
        assert updated_store["activities"][activity] is True
        assert any(
            item["ceco"] == candidate["ceco"] and item["activity"] == activity and item["valid"]
            for item in payload["submissions"]
        )
        continuation = next(item for item in payload["submissions"] if item["ceco"] == candidate["ceco"] and item["activity"] == activity)
        assert continuation["evidenceUrl"].endswith("continuacion.jpg")
        assert "Roll Out" not in {item["name"] for item in payload["activities"]}
        assert all(item["activity"] != "Roll Out" for item in payload["submissions"])
        assert payload["sources"]["cutoff"] == "2026-09-17T11:56:57"
    print("Corte histórico aprobado · conserva avance · suma cualquier actividad activa · CMS filtra actividades ocultas")


if __name__ == "__main__":
    main()
