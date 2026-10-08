#!/usr/bin/env python3
"""Regresión de Cata Cafe Oct con fuentes temporales; no modifica Forms."""
from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.build_dashboard import (
    DEFAULT_RESPONSES, DEFAULT_DIRECTORY, DEFAULT_SETTINGS, DEFAULT_CMS,
    DEFAULT_CUTOVER, build_payload, file_sha256, load_cutover,
)

ACTIVITY = "Cata Cafe Oct"


def main():
    sources = [DEFAULT_RESPONSES, DEFAULT_DIRECTORY, DEFAULT_CMS]
    before = [file_sha256(p) for p in sources]
    cut = load_cutover(DEFAULT_CUTOVER)

    def build(forms=DEFAULT_RESPONSES, cms=DEFAULT_CMS):
        return build_payload(forms, DEFAULT_DIRECTORY, DEFAULT_SETTINGS, cms,
                             baseline_path=cut["baseline"], cutoff=cut["cutoff"],
                             cutover_config_path=DEFAULT_CUTOVER)

    def activity(data):
        return next(a for a in data["activities"] if a["name"] == ACTIVITY)

    current = build()
    assert activity(current)["requireEvidence"] is True
    store = next(s for s in current["stores"] if s["ceco"] == "38368")
    assert store["activities"][ACTIVITY] is True
    matches = [s for s in current["submissions"]
               if s["ceco"] == "38368" and s["activity"] == ACTIVITY]
    assert len(matches) == 1 and matches[0]["valid"]
    assert matches[0]["evidenceSourceHeader"] == "Evidencia_Cata_Cafe_Oct"
    with tempfile.TemporaryDirectory(prefix="ops-cata-") as folder:
        folder = Path(folder)
        book = load_workbook(DEFAULT_RESPONSES)
        sheet = book.active
        headers = [str(c.value or "").strip() for c in sheet[1]]
        id_col = headers.index("Id")
        evidence_col = headers.index("Evidencia_Cata_Cafe_Oct")
        test_row = next(r for r in sheet.iter_rows(min_row=2)
                        if str(r[id_col].value) == "1851")
        assert matches[0]["evidenceUrl"] == test_row[evidence_col].value
        values = [c.value for c in test_row]
        values[id_col] = 999999
        sheet.append(values)
        duplicate = folder / "duplicate.xlsx"
        book.save(duplicate)
        repeated = build(duplicate)
        assert activity(repeated)["completedStores"] == activity(current)["completedStores"]
        assert repeated["summary"]["completedCompletions"] == current["summary"]["completedCompletions"]
        # Sin evidencia o con un dominio no autorizado, la prueba no cumple.
        for value in (None, "https://example.com/evidencia.jpg"):
            book = load_workbook(DEFAULT_RESPONSES)
            sheet = book.active
            row = next(r for r in sheet.iter_rows(min_row=2)
                       if str(r[id_col].value) == "1851")
            row[evidence_col].value = value
            forms = folder / "invalid.xlsx"
            book.save(forms)
            invalid = build(forms)
            assert not next(s for s in invalid["stores"] if s["ceco"] == "38368")["activities"][ACTIVITY]
            assert activity(invalid)["completedStores"] == activity(current)["completedStores"] - 1
        cms_book = load_workbook(DEFAULT_CMS)
        sheet = cms_book["Actividades"]
        row = next(r for r in sheet.iter_rows() if r[1].value == ACTIVITY)
        row[5].value = "No"
        cms = folder / "inactive.xlsx"
        cms_book.save(cms)
        inactive = build(cms=cms)
        assert ACTIVITY not in {a["name"] for a in inactive["activities"]}
        for previous in inactive["activities"]:
            assert previous == next(a for a in current["activities"] if a["name"] == previous["name"])
        assert current["summary"]["expectedCompletions"] - inactive["summary"]["expectedCompletions"] == activity(current)["applicableStores"]
        assert current["summary"]["completedCompletions"] - inactive["summary"]["completedCompletions"] == activity(current)["completedStores"]
    assert before == [file_sha256(p) for p in sources]
    print("Cata Cafe Oct aprobada: CeCo 38368, vínculo exacto, duplicados, evidencia obligatoria, dominio seguro, CMS inactivo y conteos previos intactos")


if __name__ == "__main__":
    main()
