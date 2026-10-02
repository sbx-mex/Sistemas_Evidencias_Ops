#!/usr/bin/env python3
"""Regresiones de cruce y CMS con fuentes temporales, sin tocar Excel operativos."""
from datetime import datetime
from pathlib import Path
import sys
import tempfile
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.build_dashboard import build_payload, find_header, load_cms, load_directory, load_settings, file_sha256


def rejected(action, message):
    try:
        action()
    except ValueError as error:
        assert message in str(error), str(error)
    else:
        raise AssertionError('Se aceptó un escenario inválido: ' + message)


def main():
    sources = [ROOT / 'cms/Sistema_Evidencias_OPS_CMS.xlsx', ROOT / 'cms/Directorio.xlsx', ROOT / 'cms/Sistema de Evidencias OPS.xlsx']
    before = [file_sha256(path) for path in sources]
    with tempfile.TemporaryDirectory(prefix='ops-cruces-') as folder:
        temp = Path(folder)
        cms, forms, baseline = (temp / name for name in ('cms.xlsx', 'forms.xlsx', 'baseline.xlsx'))
        book = load_workbook(sources[0])
        sheet = book['Actividades']
        header, cols = find_header(sheet, {'orden', 'actividad', 'descripcion', 'fecha inicio', 'fecha limite', 'activo'})
        name = 'Auditoría cruce seguro'
        for activity, order in ((name, -2), ('Auditoría inactiva', -1)):
            row = [None] * sheet.max_column
            for field, value in {'orden': order, 'actividad': activity, 'descripcion': 'Prueba', 'activo': 'Si' if activity == name else 'No', 'evidencia requerida': 'Si'}.items():
                row[cols[field]] = value
            sheet.append(row)
        book.save(cms)
        _, _, settings, _ = load_cms(cms)
        stores, _, _ = load_directory(sources[1], load_settings(ROOT / 'config/settings.json', settings), cms)
        ceco = next(iter(stores))
        other = next(value for value in stores if value != ceco)
        url = 'https://grupovips-my.sharepoint.com/prueba/valida.jpg'
        headers = ['Id', 'Hora de finalización', 'CeCo', 'CeCo1', 'Actividad', 'Evidencia', 'Correo electrónico', 'Nombre']
        def write(path, rows):
            workbook = Workbook(); ws = workbook.active; ws.append(headers)
            for row in rows: ws.append(row)
            workbook.save(path)
        def payload(**kwargs):
            return build_payload(forms, sources[1], ROOT / 'config/settings.json', cms, **kwargs)
        valid = [1, datetime(2026, 10, 1, 10), ceco, ceco + '.0', name, url, '', '']
        invalid = [2, 'fecha inválida', ceco, '', name, url.replace('valida', 'incorrecta'), '', '']
        conflict = [3, datetime(2026, 10, 1, 11), ceco, other, name, url, '', '']
        inactive = [4, 'fecha inválida', ceco, other, 'Auditoría inactiva', url, '', '']
        write(forms, [valid, invalid, conflict, inactive])
        data = payload()
        assert data['summary']['completedCompletions'] == 1
        assert len(data['submissions']) == 1 and data['submissions'][0]['evidenceUrl'] == url
        assert data['activities'][0]['name'] == name
        assert [item['order'] for item in data['activities']] == sorted(item['order'] for item in data['activities'])
        assert {tuple(item['reasons']) for item in data['quality']['quarantinedResponses']} == {('finished',), ('ceco',)}
        assert all(data['quality']['stabilityControls'].values())
        # Con corte activo, la fecha inválida no gana como supuesta fecha vacía.
        write(baseline, [valid]); write(forms, [invalid])
        cut = payload(baseline_path=baseline, cutoff=datetime(2026, 10, 1, 10))
        assert cut['summary']['completedCompletions'] == 1 and cut['submissions'][0]['evidenceUrl'] == url
        # Política estricta se respeta para conflictos y vínculos inseguros.
        config = book['Configuracion']; _, cfg = find_header(config, {'clave', 'valor'})
        policy_row = next(r for r in range(1, config.max_row + 1) if config.cell(r, cfg['clave'] + 1).value == 'responseErrorPolicy')
        config.cell(policy_row, cfg['valor'] + 1, 'Bloquear archivo'); book.save(cms)
        rejected(payload, 'Bloquear archivo')
        write(forms, [[5, datetime(2026, 10, 1, 12), ceco, '', name, 'https://example.invalid/foto.jpg', '', '']])
        rejected(payload, 'unsafe-evidence-link')
        write(forms, [inactive]); assert payload()['summary']['completedCompletions'] == 0
        # Dos columnas Estatus no pueden cambiar qué tienda entra al denominador.
        catalog = book['Tiendas Abiertas']; row, _ = find_header(catalog, {'cc', 'cc nombre', 'region', 'estatus', 'dm'})
        catalog.cell(row, catalog.max_column + 1, 'Estatus'); book.save(cms)
        rejected(lambda: load_directory(sources[1], load_settings(ROOT / 'config/settings.json', settings), cms), 'encabezados duplicados')
    assert before == [file_sha256(path) for path in sources]
    print('Cruces aprobados: CeCo/CeCo1 equivalentes y contradictorios, fecha inválida con corte, CMS orden/visibilidad, bloqueo estricto y Estatus duplicado')


if __name__ == '__main__':
    main()
