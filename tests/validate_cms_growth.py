#!/usr/bin/env python3
"""Audita ampliaciones del CMS con Excel temporales; nunca edita fuentes reales."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys
import tempfile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.build_dashboard import (
    build_payload, file_sha256, find_header, key_text, load_cms, load_directory,
    load_settings, manager_photo,
)
from scripts.validate_sources import validate_cms_engine, validate_directory_engine


def append_fields(sheet, required, values):
    _, columns = find_header(sheet, required)
    row = [None] * sheet.max_column
    for field, value in values.items():
        row[columns[field]] = value
    sheet.append(row)


def rejected(action, message):
    try:
        action()
    except ValueError as error:
        assert message in str(error), str(error)
    else:
        raise AssertionError("Se aceptó un escenario inválido: " + message)


def main():
    sources = [ROOT / "cms/Sistema_Evidencias_OPS_CMS.xlsx", ROOT / "cms/Directorio.xlsx", ROOT / "cms/Sistema de Evidencias OPS.xlsx"]
    before = [file_sha256(path) for path in sources]
    activities, managers, cms_settings, _ = load_cms(sources[0])
    settings = load_settings(ROOT / "config/settings.json", cms_settings)
    stores, _, _ = load_directory(sources[1], settings)
    original_regionals = cms_settings["_organization"]["regionalDirectors"]
    with tempfile.TemporaryDirectory(prefix="ops-crecimiento-") as folder:
        temp = Path(folder)
        cms, directory, forms = temp / "cms.xlsx", temp / "directory.xlsx", temp / "forms.xlsx"
        book = load_workbook(sources[0])
        directory_book = load_workbook(sources[1])
        directory_sheet = directory_book[settings["directorySheet"]]
        manager_contract = {"dm", "nombre corto", "foto webp", "activo"}
        org_contract = {"nivel", "region", "nombre", "rol", "foto webp", "activo", "orden"}
        store_contract = {"cc", "cc nombre", "region", "estatus", "dm"}
        new_cecos = [str(60000 + index) for index in range(240)]
        assert not set(new_cecos).intersection(stores)
        for index in range(60):
            region = f"Región de prueba {index // 6 + 1:02d}"
            dm = f"Responsable de prueba {index + 1:02d}"
            append_fields(book["Gerentes"], manager_contract, {
                "dm": dm, "nombre corto": dm, "foto webp": f"assets/dm/prueba-{index}.webp", "activo": "Si",
            })
            for offset in range(4):
                values = {"cc": new_cecos[index * 4 + offset], "cc nombre": f"Tienda de prueba {index * 4 + offset + 1}", "region": region, "estatus": "Abierta", "dm": dm}
                append_fields(directory_sheet, store_contract, values)
                append_fields(book["Tiendas Abiertas"], store_contract, values)
        for index in range(10):
            append_fields(book["Organigrama"], org_contract, {
                "nivel": 2, "region": f"Región de prueba {index + 1:02d}", "nombre": f"Directora de prueba {index + 1:02d}", "rol": "Director Regional", "foto webp": f"assets/director/prueba-{index}.webp", "activo": "Si", "orden": 100 + index,
            })
        activity_contract = {"orden", "actividad", "descripcion", "fecha inicio", "fecha limite", "activo"}
        for index in range(2):
            append_fields(book["Actividades"], activity_contract, {
                "orden": 10000 + index, "actividad": f"Actividad de prueba {index + 1}", "descripcion": "Auditoría temporal", "fecha inicio": None, "fecha limite": None, "activo": "Si", "evidencia requerida": "Si", "prioridad": "Media", "estado fecha": "Sin fecha",
            })
        book.save(cms)
        directory_book.save(directory)
        forms_book = Workbook()
        sheet = forms_book.active
        sheet.append(["Id", "Hora de inicio", "Hora de finalización", "Correo electrónico", "Nombre", "CeCo", "Actividad", "Evidencia"])
        latest_link = "https://grupovips-my.sharepoint.com/auditoria/reciente.jpg"
        old_link = "https://grupovips-my.sharepoint.com/auditoria/anterior.jpg"
        # La respuesta reciente va antes de la anterior: la fecha decide.
        sheet.append([2, datetime(2026, 10, 1, 11), datetime(2026, 10, 1, 12), "tienda@alsea.com.mx", "Tienda", new_cecos[0], "Actividad de prueba 1", latest_link])
        sheet.append([1, datetime(2026, 10, 1, 9), datetime(2026, 10, 1, 10), "gerente@alsea.com.mx", "Gerente", new_cecos[0], "Actividad de prueba 1", old_link])
        sheet.append([3, datetime(2026, 10, 1, 12), datetime(2026, 10, 1, 13), "gerente@alsea.com.mx", "Gerente", "99999", "Actividad de prueba 1", latest_link])
        forms_book.save(forms)
        audit = validate_cms_engine(cms)
        assert audit["managers"] == len(managers) + 60
        assert audit["regionalDirectors"] == len(original_regionals) + 10
        assert validate_directory_engine(directory, cms, ROOT / "config/settings.json")["stores"] == len(stores) + 240
        payload = build_payload(forms, directory, ROOT / "config/settings.json", cms)
        assert payload["summary"]["stores"] == len(stores) + 240
        assert payload["summary"]["activities"] == len(activities) + 2
        assert len(payload["dms"]) == len({store["dm"] for store in stores.values()}) + 60
        assert len(payload["organization"]["regionalDirectors"]) == len(original_regionals) + 10
        assert payload["summary"]["completedCompletions"] == 1
        selected = [item for item in payload["submissions"] if item["ceco"] == new_cecos[0]]
        assert len(selected) == 1 and selected[0]["evidenceUrl"] == latest_link
        assert payload["quality"]["unknownCeCos"] == ["99999"]
        assert all(not item["photo"] and item["photoStatus"] == "Pendiente" for item in payload["dms"] if item["dm"].startswith("Responsable de prueba"))
        assert sum(item["stores"] for item in payload["organization"]["regionalDirectors"] if item["region"].startswith("Región de prueba")) == 240
        assert manager_photo("Nuevo DM", "Nuevo DM", "assets/dm/no-existe.webp") == ("", "Pendiente")
        rejected(lambda: manager_photo("DM", "DM", "../insegura.webp"), "sale del proyecto")
        rejected(lambda: manager_photo("DM", "DM", "https://example.com/foto.webp"), "Ruta WebP inválida")
        # Duplicados de organización siguen bloqueando, aun si faltan fotografías.
        append_fields(book["Organigrama"], org_contract, {
            "nivel": 2, "region": "Región de prueba 01", "nombre": "Otro director", "rol": "Director Regional", "foto webp": None, "activo": "Si", "orden": 999,
        })
        book.save(cms)
        rejected(lambda: load_cms(cms), "duplicados")
        rejected(lambda: validate_cms_engine(cms), "duplicados")
    assert before == [file_sha256(path) for path in sources], "La auditoría alteró fuentes reales"
    print("Crecimiento CMS aprobado · +240 tiendas · +60 DM · +10 RD · +2 actividades · fotos pendientes · última encuesta por CeCo/actividad · CeCo desconocido aislado · duplicados bloqueados")


if __name__ == "__main__":
    main()
