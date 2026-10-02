#!/usr/bin/env python3
"""Valida altas, bajas y reasignaciones sin modificar los catálogos reales."""
from __future__ import annotations

from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.build_dashboard import file_sha256, find_header, load_cms, load_directory, load_settings


def reject(action, text):
    try:
        action()
    except ValueError as error:
        assert text in str(error), str(error)
    else:
        raise AssertionError("Se aceptó un catálogo inválido: " + text)


def main():
    cms = ROOT / "cms/Sistema_Evidencias_OPS_CMS.xlsx"
    directory = ROOT / "cms/Directorio.xlsx"
    before = file_sha256(cms), file_sha256(directory)
    _, _, config, _ = load_cms(cms)
    settings = load_settings(ROOT / "config/settings.json", config)
    settings["storeCatalogSource"] = "CMS"
    current, sheet_name, status = load_directory(directory, settings, cms)
    reference, _, _ = load_directory(directory, {**settings, "storeCatalogSource": "Directorio"}, cms)
    assert sheet_name == "Tiendas Abiertas" and status["catalogSource"] == "CMS"
    assert status["newCeCos"] == sorted(set(current) - set(reference))
    with tempfile.TemporaryDirectory(prefix="ops-catalogo-") as folder:
        path = Path(folder) / "cms.xlsx"
        book = load_workbook(cms)
        sheet = book["Tiendas Abiertas"]
        header, cols = find_header(sheet, {"cc", "cc nombre", "region", "estatus", "dm"})
        first, second = header + 1, header + 2
        changed_ceco = str(sheet.cell(first, cols["cc"] + 1).value).strip()
        removed_ceco = str(sheet.cell(second, cols["cc"] + 1).value).strip()
        sheet.cell(first, cols["cc nombre"] + 1, "Tienda con nombre actualizado")
        sheet.cell(first, cols["dm"] + 1, "Responsable actualizado")
        sheet.delete_rows(second)
        new_ceco = next(str(value) for value in range(61000, 62000) if str(value) not in current)
        row = [None] * sheet.max_column
        for key, value in {"cc": new_ceco, "cc nombre": "Nueva tienda", "region": next(iter(current.values()))["region"], "estatus": "Abierta", "dm": "Responsable nuevo"}.items():
            row[cols[key]] = value
        sheet.append(row)
        book.save(path)
        updated, _, audit = load_directory(directory, settings, path)
        assert len(updated) == len(current)
        assert new_ceco in updated and removed_ceco not in updated
        assert updated[changed_ceco]["store"] == "Tienda con nombre actualizado"
        assert updated[changed_ceco]["dm"] == "Responsable actualizado"
        assert new_ceco in audit["newCeCos"]
        # El Directorio antiguo sigue conteniendo la tienda retirada; no la revive.
        if removed_ceco in reference:
            assert removed_ceco in audit["retiredCeCos"]
        if changed_ceco in reference:
            assert changed_ceco in audit["reassignedCeCos"]
        sheet.append(row)
        book.save(path)
        reject(lambda: load_directory(directory, settings, path), "CeCo duplicado")
        sheet.delete_rows(sheet.max_row)
        sheet.cell(sheet.max_row, cols["cc"] + 1, "CeCo incorrecto")
        book.save(path)
        reject(lambda: load_directory(directory, settings, path), "CeCo inválido")
        sheet.cell(sheet.max_row, cols["cc"] + 1, new_ceco)
        sheet.cell(sheet.max_row, cols["region"] + 1).value = None
        book.save(path)
        reject(lambda: load_directory(directory, settings, path), "Tienda CMS incompleta")
        reject(lambda: load_directory(directory, {**settings, "storeCatalogSource": "Mezclar"}, path), "sólo acepta")
    assert before == (file_sha256(cms), file_sha256(directory))
    print(f"Catálogo vigente aprobado · CMS {len(current)} tiendas · respaldo {len(reference)} · altas / bajas / reasignaciones · CeCo y duplicados bloqueados · fuentes sin cambios")


if __name__ == "__main__":
    main()
