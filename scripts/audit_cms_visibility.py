#!/usr/bin/env python3
"""Audita con fuentes reales y copias temporales que sólo el CMS autorice actividades."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import timedelta
import json
from pathlib import Path
import sys
import tempfile

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook, load_workbook
from scripts.build_dashboard import (
    DEFAULT_CMS, DEFAULT_CUTOVER, DEFAULT_DIRECTORY, DEFAULT_RESPONSES,
    DEFAULT_SETTINGS, active_activity_catalog, build_payload,
    canonical_cms_activity, clean_text, compact_key, file_sha256,
    find_header, find_response_source, is_yes, load_cms, load_cutover, load_responses,
)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def operational(data: dict) -> dict:
    return {key: data[key] for key in (
        "summary", "activities", "stores", "dms", "attention", "submissions",
        "quantityModules", "surveyModules", "lastUpdated", "lastUpdatedDisplay",
    )}


def check_catalog(data: dict, authorized: set[str]) -> None:
    names = [item["name"] for item in data["activities"]]
    require(len(names) == len(authorized) and set(names) == authorized,
            "El tablero no coincide exactamente con las actividades activas CMS")
    require(data["summary"]["activities"] == len(authorized), "Denominador de actividades incorrecto")
    for store in data["stores"]:
        require(set(store["activities"]) == authorized, "La matriz de tienda contiene una actividad omitida")
        require(set(store["applicableActivities"]) == authorized, "Aplicabilidad contiene una actividad omitida")
    for group in ("submissions", "quantityModules", "surveyModules"):
        require(all(item["activity"] in authorized for item in data[group]),
                f"Una actividad omitida aparece en {group}")
    expected = sum(item["applicableStores"] for item in data["activities"])
    completed = sum(item["completedStores"] for item in data["activities"])
    require(data["summary"]["expectedCompletions"] == expected, "Ideal no reconciliado")
    require(data["summary"]["completedCompletions"] == completed, "Cumplimientos no reconciliados")
    require(data["summary"]["pendingCompletions"] == expected - completed, "Pendientes no reconciliados")


def blank_excluded(source: Path, target: Path, active_keys: set[str]) -> None:
    responses, schema = load_responses(source)
    book = load_workbook(source)
    sheet = book[schema["sheet"]]
    for response in responses:
        if compact_key(response["activity"]) not in active_keys:
            for cell in sheet[response["row"]]:
                cell.value = None
    book.save(target)
    book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, help="Guardar resultados JSON")
    args = parser.parse_args()
    cut = load_cutover(DEFAULT_CUTOVER)
    require(bool(cut), "Se requiere el corte configurado para auditar ambas fuentes")
    sources = [DEFAULT_CMS, DEFAULT_RESPONSES, DEFAULT_DIRECTORY, DEFAULT_SETTINGS,
               DEFAULT_CUTOVER, cut["baseline"]]
    before = {str(path.relative_to(ROOT)): file_sha256(path) for path in sources}
    activities, _, settings, _ = load_cms(DEFAULT_CMS)
    active = {item["name"] for item in activities}
    excluded = settings["_excludedActivityNames"]
    active_keys = {compact_key(name) for name in active}
    catalog = active_activity_catalog(activities)
    require(all(canonical_cms_activity(name, *catalog, excluded) is None for name in excluded),
            "Una actividad Activo=No fue autorizada")
    checks = []

    def approve(label: str) -> None:
        checks.append(label)

    def build(forms=DEFAULT_RESPONSES, cms=DEFAULT_CMS, baseline=None, with_cut=True):
        return build_payload(forms, DEFAULT_DIRECTORY, DEFAULT_SETTINGS, cms,
                             baseline_path=(baseline or cut["baseline"]) if with_cut else None,
                             cutoff=cut["cutoff"] if with_cut else None)

    current = build()
    check_catalog(current, active)
    approve("Catálogo, tiendas, módulos y conteos sólo contienen actividades activas CMS")
    source_stats = []
    for path in (cut["baseline"], DEFAULT_RESPONSES):
        rows, _ = load_responses(path)
        outside = Counter(clean_text(row["activity"]) for row in rows
                          if compact_key(row["activity"]) not in active_keys)
        source_stats.append({"archivo": path.name, "filasLeidas": len(rows),
                             "filasFueraCMSActivo": sum(outside.values()),
                             "actividadesOmitidas": dict(sorted(outside.items()))})
        check_catalog(build(forms=path, with_cut=False), active)
    approve("Ambos Excel respetan el CMS también sin unión del corte")

    with tempfile.TemporaryDirectory(prefix="ops-cms-audit-") as directory:
        temp = Path(directory)
        forms_clean, baseline_clean = temp / "forms.xlsx", temp / "baseline.xlsx"
        blank_excluded(DEFAULT_RESPONSES, forms_clean, active_keys)
        blank_excluded(cut["baseline"], baseline_clean, active_keys)
        require(operational(build(forms=forms_clean, baseline=baseline_clean)) == operational(current),
                "Retirar filas fuera del CMS activo cambió el resultado operativo")
        approve("Conservar u omitir historia inactiva no cambia avance, evidencia ni fecha")

        for item in activities:
            book = load_workbook(DEFAULT_CMS)
            sheet = book["Actividades"]
            header, cols = find_header(sheet, {"actividad", "activo"})
            for row in range(header + 1, sheet.max_row + 1):
                if clean_text(sheet.cell(row, cols["actividad"] + 1).value) == item["name"]:
                    sheet.cell(row, cols["activo"] + 1, "No")
            target = temp / "cms-off.xlsx"
            book.save(target)
            book.close()
            off = build(cms=target)
            check_catalog(off, active - {item["name"]})
            approve(f"Activo=No retira historia, módulo y denominador: {item['name']}")

        # Un nombre externo muy parecido antes se aceptaba por afinidad.
        chosen = activities[-1]["name"]
        outsider = chosen + " POS"
        require(canonical_cms_activity(outsider, *catalog) is None,
                "Un nombre externo parecido fue autorizado por afinidad")
        require(canonical_cms_activity(chosen.upper().replace(" ", "_"), *catalog) == chosen,
                "La normalización de formato perdió una actividad autorizada")
        book = load_workbook(DEFAULT_CMS)
        sheet = book["Actividades"]
        header, cols = find_header(sheet, {"actividad", "activo"})
        row = sheet.max_row + 1
        sheet.cell(row, cols["actividad"] + 1, outsider)
        sheet.cell(row, cols["activo"] + 1, "No")
        similar_cms = temp / "similar-cms.xlsx"
        book.save(similar_cms)
        book.close()
        # La fila externa y su encabezado no pueden acreditar la actividad activa.
        forms = temp / "similar.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["Id", "Hora de finalización", "CeCo",
                      "Selecciona la actividad que deseas registrar",
                      "Evidencia_" + outsider, "Evidencia_" + chosen])
        timestamp = cut["cutoff"] + timedelta(days=30)
        sheet.append([1, timestamp, current["stores"][0]["ceco"], outsider,
                      "https://grupovips-my.sharepoint.com/evidencias/omitida.jpg", ""])
        sheet.append([2, timestamp, current["stores"][0]["ceco"], chosen,
                      "https://grupovips-my.sharepoint.com/evidencias/columna-incorrecta.jpg", ""])
        book.save(forms)
        book.close()
        isolated = build(forms=forms, cms=similar_cms, with_cut=False)
        require(isolated["summary"]["completedCompletions"] == 0,
                "Un nombre o encabezado inactivo acreditó una actividad activa")
        require(outsider in isolated["quality"]["hiddenActivities"], "No se registró la omisión")
        approve("Nombre externo similar y encabezado de actividad No no acreditan cumplimiento")

        # Las contradicciones de identidad se bloquean, no se resuelven por orden.
        book = load_workbook(DEFAULT_CMS)
        sheet = book["Actividades"]
        row = sheet.max_row + 1
        sheet.cell(row, cols["actividad"] + 1, chosen)
        sheet.cell(row, cols["activo"] + 1, "No")
        conflict = temp / "conflict.xlsx"
        book.save(conflict)
        book.close()
        try:
            load_cms(conflict)
        except ValueError as error:
            require("activa e inactiva" in str(error), "Error inesperado de identidad CMS")
        else:
            raise RuntimeError("CMS contradictorio no fue rechazado")
        approve("Actividad duplicada con Si/No contradictorios bloquea la carga")

    # Exportaciones vigentes deben llevar exactamente el catálogo activo.
    export = load_workbook(ROOT / "exports/Resumen_Evidencias_OPS.xlsx", read_only=True, data_only=True)
    exported = [row[1] for row in export["Actividades"].iter_rows(min_row=5, values_only=True) if row[1]]
    require(len(exported) == len(active) and set(exported) == active, "Excel exportado desactualizado")
    require(all(row[2] in active for row in export["Evidencias"].iter_rows(min_row=5, values_only=True)),
            "Excel exportado incluye evidencia inactiva")
    export.close()
    approve("Excel exportado sólo incluye actividades y evidencias activas")
    require(before == {str(path.relative_to(ROOT)): file_sha256(path) for path in sources},
            "La auditoría modificó una fuente")
    approve("CMS, Directorio, ambos Forms y configuración conservan sus SHA256")
    report = {"estado": "APROBADO", "controles": len(checks), "verificaciones": checks,
              "activas": [item["name"] for item in activities], "omitidasCMS": excluded,
              "fuentes": source_stats, "sha256Fuentes": before,
              "resumen": current["summary"], "ultimaRespuesta": current["lastUpdatedDisplay"],
              "cecosAislados": current["quality"].get("unknownCeCos", [])}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"CMS seguro aprobado · {len(checks)}/{len(checks)} controles · {len(active)} activas · {len(excluded)} omitidas")


if __name__ == "__main__":
    main()
