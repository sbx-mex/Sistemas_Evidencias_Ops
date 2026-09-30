#!/usr/bin/env python3
"""Valida evidencias múltiples, aislamiento y totales de todos los tipos publicados."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta
import json
from pathlib import Path
import sys
import tempfile
import argparse

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook, load_workbook
from scripts.build_dashboard import (
    DEFAULT_CMS, DEFAULT_CUTOVER, DEFAULT_DIRECTORY, DEFAULT_RESPONSES, DEFAULT_SETTINGS,
    MULTI_EVIDENCE_CONFIG, build_payload, clean_text, compact_key, file_sha256,
    find_header, find_response_source, load_cms, load_cutover, load_responses,
    matching_columns, normalize_allowed_hosts,
    recover_response_ceco,
)
from scripts.export_excel import build_workbook


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    activities, _, settings, _ = load_cms(DEFAULT_CMS)
    names = [item["name"] for item in activities]
    cut = load_cutover(DEFAULT_CUTOVER)
    sources = [DEFAULT_CMS, DEFAULT_DIRECTORY, DEFAULT_RESPONSES, DEFAULT_SETTINGS, DEFAULT_CUTOVER]
    if cut:
        sources.append(cut["baseline"])
    hashes = {p.relative_to(ROOT).as_posix(): file_sha256(p) for p in sources}
    published = json.loads((ROOT / "data/dashboard.json").read_text(encoding="utf-8"))
    payload = build_payload(DEFAULT_RESPONSES, DEFAULT_DIRECTORY, DEFAULT_SETTINGS, DEFAULT_CMS,
                            baseline_path=cut["baseline"] if cut else None,
                            cutoff=cut["cutoff"] if cut else None,
                            cutover_config_path=DEFAULT_CUTOVER if cut else None)
    require(all(published[key] == payload[key] for key in (
        "activities", "submissions", "summary", "quantityModules", "surveyModules", "lastUpdated")),
        "El tablero guardado no corresponde a las respuestas actuales")
    controls = []

    def approve(label: str) -> None:
        controls.append(label)

    approve("Tablero sincronizado con CMS y ambas fuentes Forms")
    rows, schema = load_responses(DEFAULT_RESPONSES, names, settings.get("_quantityConfig"),
                                  settings.get("_excludedActivityNames"))
    row_counts = Counter(row["activity"] for row in rows)
    valid_counts = Counter(row["activity"] for row in payload["submissions"] if row["valid"])
    require(all(name in names for name in valid_counts), "Se publicó una actividad fuera del CMS")
    require(len({(r['ceco'], r['activity']) for r in payload['submissions']}) == len(payload['submissions']),
            "Se duplicó una respuesta de tienda/actividad")
    approve("Una sola respuesta ganadora por tienda/actividad, dentro del CMS")

    # Contraste con celdas originales, sin usar el resolvedor de evidencias.
    # Se comprueba que cada pareja publicada venga de la misma fila Forms.
    stores = {store["ceco"]: store for store in payload["stores"]}
    raw_profiles = {key: set() for key in MULTI_EVIDENCE_CONFIG}
    for source in ([cut["baseline"]] if cut else []) + [DEFAULT_RESPONSES]:
        parsed, _ = load_responses(source, names, settings.get("_quantityConfig"),
                                   settings.get("_excludedActivityNames"))
        parsed_by_row = {r["row"]: r for r in parsed}
        book = load_workbook(source, read_only=True, data_only=True)
        sheet, header_row, headers = find_response_source(book, names)
        for key, profile in MULTI_EVIDENCE_CONFIG.items():
            header_groups = [matching_columns(headers, stage["headers"]) for stage in profile["stages"]]
            for row_number, raw in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                source_row = parsed_by_row.get(row_number)
                if not source_row or compact_key(source_row["activity"]) != key or not all(header_groups):
                    continue
                resolved_ceco, _ = recover_response_ceco(source_row, stores,
                    enabled=bool(settings.get("trustedCeCoRecovery", True)))
                stage_indices = [next((i for i in group if clean_text(raw[i])), group[0])
                                 for group in header_groups]
                files = tuple((stage["label"], clean_text(headers[index]), clean_text(raw[index]))
                              for stage, index in zip(profile["stages"], stage_indices))
                raw_profiles[key].add((resolved_ceco, files))
        book.close()
    for record in payload["submissions"]:
        key = compact_key(record["activity"])
        if key not in MULTI_EVIDENCE_CONFIG:
            continue
        expected_labels = [stage["label"] for stage in MULTI_EVIDENCE_CONFIG[key]["stages"]]
        files = record.get("evidenceFiles", [])
        require([f["label"] for f in files] == expected_labels, "Falta una etapa de evidencia múltiple")
        signature = (record["ceco"], tuple((f["label"], f["sourceHeader"], f["url"]) for f in files))
        require(signature in raw_profiles[key], "Las evidencias múltiples no vienen de la misma fila original")
    approve("POS y Showcase publicados conservan ambos vínculos de la misma fila original")

    profile_key = "tartamaricu"
    profile = MULTI_EVIDENCE_CONFIG[profile_key]
    allowed = normalize_allowed_hosts(settings.get("evidenceAllowedHosts", "grupovips-my.sharepoint.com"))
    tarta_rows = [r for r in rows if compact_key(r["activity"]) == profile_key]
    if profile["activity"] in names:
        require(all(r.get("evidenceIssue") is None for r in tarta_rows
                    if len(r.get("evidenceFiles", [])) == 2),
                "Una pareja completa fue aislada por no reconocer sus etapas")
        require(all(schema["evidenceHeaderMap"].get(clean_text(stage["headers"][0])) == profile_key
                    for stage in profile["stages"] if tarta_rows), "Faltan encabezados de Tarta en el contrato")
    approve("Encabezados nuevos de Tarta reconocidos expresamente")

    for module in payload["quantityModules"]:
        records = [r for r in payload["submissions"] if r["activity"] == module["activity"] and r["valid"]]
        require(module["answeredStores"] == len(records), "Número de tiendas numéricas incongruente")
        for metric in module["metrics"]:
            key = metric["key"]
            total = sum(r["quantities"][key] for r in records)
            require(module["totals"][key] == total, "Total numérico incongruente")
            for grouping in ("byRegion", "byPortfolio"):
                require(sum(group["totals"][key] for group in module[grouping]) == total,
                        "Desglose por región/DM no reconcilia")
            require(all(metric["minimum"] <= r["quantities"][key] <= metric["maximum"] for r in records),
                    "Se publicaron cantidades fuera de rango")
        total = sum(module["totals"][m["key"]] for m in module["metrics"])
        require(module["totals"]["total"] == total, "Total consolidado incongruente")
        numerator = module.get("percentageMetric")
        if numerator:
            percentage = round(module["totals"][numerator] / total * 100, 1) if total else None
            require(module["totals"]["percentage"] == percentage, "Porcentaje usa un promedio incorrecto")
    approve("Jarras, FHW y Va X Cuenta reconcilian por respuesta, región y DM; porcentaje ponderado")

    # Exportador existente: todas las etapas aparecen con el vínculo intacto.
    workbook = build_workbook(payload)
    actual = sorted((row[0], row[2], row[3], row[5]) for row in
                    workbook["Evidencias"].iter_rows(min_row=5, values_only=True))
    expected = sorted((record["ceco"], record["activity"], file["label"], file["url"])
                      for record in payload["submissions"] if record.get("valid") and record.get("evidenceUrl")
                      for file in (record.get("evidenceFiles") or [{"label": "Evidencia", "url": record["evidenceUrl"]}]))
    require(actual == expected, "Excel cambia u omite una etapa o vínculo de evidencia")
    workbook.close()
    approve("Exportación Excel conserva todas las etapas y vínculos exactamente")

    # Casos adversos únicamente en copias temporales, con un CMS de prueba activo.
    with tempfile.TemporaryDirectory(prefix="ops-response-types-") as directory:
        temp = Path(directory)
        cms = temp / "cms.xlsx"
        book = load_workbook(DEFAULT_CMS)
        sheet = book["Actividades"]
        header, cols = find_header(sheet, {"actividad", "activo"})
        for row in range(header + 1, sheet.max_row + 1):
            if compact_key(sheet.cell(row, cols["actividad"] + 1).value) == profile_key:
                sheet.cell(row, cols["activo"] + 1, "Si")
                break
        else:
            row = sheet.max_row + 1
            sheet.cell(row, cols["actividad"] + 1, profile["activity"])
            sheet.cell(row, cols["activo"] + 1, "Si")
        book.save(cms)
        book.close()
        ceco = payload["stores"][0]["ceco"]
        ordinary = next(a["name"] for a in activities
                        if compact_key(a["name"]) not in MULTI_EVIDENCE_CONFIG
                        and a["name"] not in {m["activity"] for m in payload["quantityModules"]})
        timestamp = datetime(2026, 9, 29, 19)
        pos, showcase = [stage["headers"][0] for stage in profile["stages"]]
        url = "https://" + sorted(allowed)[0] + "/evidencias/"
        base = ["Id", "Hora de finalización", "CeCo", "Actividad"]

        def simulate(extra_headers, extra_values, activity=profile["activity"]):
            source = temp / "forms.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(base + extra_headers)
            for i, values in enumerate(extra_values, 1):
                sheet.append([i, timestamp + timedelta(seconds=i), ceco, activity, *values])
            book.save(source)
            book.close()
            return build_payload(source, DEFAULT_DIRECTORY, DEFAULT_SETTINGS, cms)

        pair = simulate([showcase + "\n", pos + "\n"], [[url + "showcase.jpg", url + "pos.jpg"]])
        require(pair["summary"]["completedCompletions"] == 1, "Pareja válida no completa exactamente una actividad")
        require([f["label"] for f in pair["submissions"][0]["evidenceFiles"]] == ["POS", "Showcase"],
                "Reordenar columnas cambió el orden de etapas")
        approve("Pareja completa: una actividad y dos archivos, con columnas reordenadas")

        repeated = simulate([pos, showcase], [[url + "old-pos.jpg", url + "old-showcase.jpg"],
                                             [url + "new-pos.jpg", url + "new-showcase.jpg"]])
        require(repeated["summary"]["completedCompletions"] == 1, "Se duplicó un envío repetido")
        require([f["url"] for f in repeated["submissions"][0]["evidenceFiles"]] ==
                [url + "new-pos.jpg", url + "new-showcase.jpg"], "Se mezclaron archivos de dos respuestas")
        approve("Reenvío más reciente reemplaza la pareja completa sin duplicar ni mezclar")

        missing = simulate([pos, showcase], [[url + "pos.jpg", ""]])
        missing_header = simulate([pos], [[url + "pos.jpg"]])
        require(all(x["summary"]["completedCompletions"] == 0 for x in (missing, missing_header)),
                "Una etapa o encabezado faltante acreditó Tarta")
        approve("Archivo o encabezado de Showcase faltante deja la respuesta fuera de cumplimiento")

        duplicates = simulate([pos, pos + " (2)", showcase], [[url + "pos.jpg", url + "pos.jpg", url + "showcase.jpg"]])
        require(duplicates["summary"]["completedCompletions"] == 1, "Duplicados equivalentes no se consolidaron")
        conflict = simulate([pos, pos + " (2)", showcase], [[url + "pos.jpg", url + "other-pos.jpg", url + "showcase.jpg"]])
        require(conflict["summary"]["completedCompletions"] == 0, "Duplicados contradictorios fueron aceptados")
        approve("Duplicados idénticos se consolidan; duplicados contradictorios se aíslan")

        mixed = simulate([pos, showcase, "Evidencia_" + ordinary],
                         [[url + "pos.jpg", url + "showcase.jpg", url + "foreign.jpg"]])
        wrong_activity = simulate([pos, showcase], [[url + "pos.jpg", url + "showcase.jpg"]], ordinary)
        require(all(x["summary"]["completedCompletions"] == 0 for x in (mixed, wrong_activity)),
                "Archivos de otra actividad acreditaron cumplimiento")
        approve("Archivos extra o actividad incorrecta se aíslan sin reasignar evidencias")

        unsafe = simulate([pos, showcase], [[url + "pos.jpg", "https://example.com/showcase.jpg"]])
        require(unsafe["summary"]["completedCompletions"] == 0 and
                unsafe["quality"]["unsafeEvidenceRows"] == [2], "Un segundo vínculo inseguro fue publicado")
        approve("Cada vínculo supera la validación HTTPS y de dominio autorizado")

    require(hashes == {p.relative_to(ROOT).as_posix(): file_sha256(p) for p in sources},
            "La auditoría modificó fuentes Excel o configuración")
    approve("Fuentes originales conservan sus SHA256")
    report = {"estado": "APROBADO", "controles": len(controls), "verificaciones": controls,
              "sha256Fuentes": hashes, "resumen": payload["summary"],
              "ultimaRespuesta": payload["lastUpdatedDisplay"],
              "filasFormsActual": len(rows), "filasBase": payload["quality"].get("cutover", {}).get("baselineRowsRead", 0),
              "porActividad": [{"actividad": name, "filasFormsActual": row_counts[name],
                                "tiendasConCumplimiento": valid_counts[name]} for name in names],
              "totalesNumericos": [{"actividad": m["activity"], "tiendas": m["answeredStores"],
                                    "totales": m["totals"]} for m in payload["quantityModules"]],
              "cecosAislados": payload["quality"]["unknownCeCos"],
              "respuestasAisladas": payload["quality"]["quarantinedResponses"]}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Tipos de respuesta aprobados · {len(controls)}/{len(controls)} controles · "
          f"{len(rows)} respuestas Forms · {len(tarta_rows)} respuestas Tarta leídas · "
          f"{valid_counts[profile['activity']]} tiendas Tarta válidas")


if __name__ == "__main__":
    main()
